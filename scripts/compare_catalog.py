#!/usr/bin/env python3
"""Compare a generated catalog against a real source file, full population.

Checks header list/order, row population (by SKU column), cell value AND
cell type (xlsx read with data_only=False so uncached formulas are visible),
and for csv/txt a raw byte-level comparison of framing (BOM, line endings).

    python3 scripts/compare_catalog.py <real> <generated> [--sku "Part Number"] [--max-diffs 5]
"""
import argparse
import csv
import io
from collections import defaultdict

import openpyxl


def _typed(v):
    if v is None:
        return ("", "blank")
    if isinstance(v, str) and v.startswith("="):
        return (v, "formula")
    return (str(v), type(v).__name__)


def sheet_names(path):
    if not path.lower().endswith((".xlsx", ".xlsm")):
        return None
    wb = openpyxl.load_workbook(path, read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def _raw_rows(path, sheet=None, delimiter=None, sku=None):
    """All non-empty rows as lists of (value, type_name) cells. For a
    multi-sheet workbook with no `sheet` given, uses the first sheet whose
    first 60 rows contain the `sku` header (e.g. Wageningen's 'Catalogus')."""
    if path.lower().endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        if not sheet and sku and len(wb.worksheets) > 1:
            for cand in wb.worksheets:
                if any(sku in r for r in cand.iter_rows(values_only=True, max_row=60)):
                    ws = cand
                    break
        rows = [[_typed(v) for v in r] for r in ws.iter_rows(values_only=True)
                if r is not None and not all(v is None for v in r)]
        wb.close()
        return rows
    text = open(path, "rb").read().decode("utf-8-sig", errors="replace")
    if text.startswith("CIF_I_V3.0"):
        delimiter = ","
    if not delimiter:
        try:
            delimiter = csv.Sniffer().sniff(text[:5000], delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ","
    return [[(v, "str" if v else "blank") for v in r]
            for r in csv.reader(io.StringIO(text), delimiter=delimiter) if r]


def read_table(path, sheet=None, sku=None, header=None, delimiter=None):
    """Return (header, rows) with each cell as (value, type_name).

    Layouts handled: header in row 1 (default); a preamble before the header
    row (the header is the first row, within the first 60, containing `sku`);
    CIF_I_V3.0 (.txt or .xlsx: metadata lines, 'FIELDNAMES: ...' header, rows
    between 'DATA' and 'ENDOFDATA'); and headerless files, where the caller
    passes the column names as `header`.
    """
    rows = _raw_rows(path, sheet, delimiter, sku)
    if not rows:
        return (header or []), []
    first = rows[0][0][0] if rows[0] else ""
    if first == "CIF_I_V3.0":
        hi = next(i for i, r in enumerate(rows) if r and r[0][0].startswith("FIELDNAMES:"))
        head = [c[0] for c in rows[hi]]
        head[0] = head[0][len("FIELDNAMES:"):].strip()
        di = next(i for i, r in enumerate(rows) if r and r[0][0] == "DATA")
        data = []
        for r in rows[di + 1:]:
            if r and r[0][0] == "ENDOFDATA":
                break
            data.append(r)
        return head, data
    if header is not None:
        return list(header), rows
    hi = 0
    if sku:
        hi = next((i for i, r in enumerate(rows[:60]) if sku in [c[0] for c in r]), 0)
    return [c[0] for c in rows[hi]], rows[hi + 1:]


def byte_framing(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        return None
    raw = open(path, "rb").read()
    return {
        "bom": raw.startswith(b"\xef\xbb\xbf"),
        # Line-ending style, not counts: embedded newlines inside quoted
        # fields change the counts without changing the file's framing.
        "line_endings": ("crlf" if raw.count(b"\r\n") == raw.count(b"\n") else
                         "lf" if b"\r\n" not in raw else "mixed"),
        "leading_blank_line": raw.lstrip(b"\xef\xbb\xbf")[:2] in (b"\r\n", b"\n\n") or raw.lstrip(b"\xef\xbb\xbf")[:1] == b"\n",
        "ends_with_newline": raw.endswith(b"\n"),
        "tail": raw[-40:],
    }


def compare(real_path, generated_path, sku=None, header=None, delimiter=None, sheet=None):
    """Full-population comparison. Returns a dict:
    {real_cols, gen_cols, real_rows, gen_rows, header_diffs: [(col, real, gen)],
     framing: None | {real, gen}, sku_missing: bool, only_real, only_gen,
     common, value_diffs: {label: [(sku, real, gen)]},
     type_diffs: {label: {(real_type, gen_type): n}},
     type_diff_skus: {label: {(real_type, gen_type): [sku, ...]}}, ok}.
    `header` names the columns of headerless files (csv_format.header_row
    false); `delimiter` overrides CSV sniffing. `sheet` (the template's
    sheet_name) is compared when BOTH workbooks have it - needed when a
    leading reference sheet also contains the SKU header (sanquin)."""
    if sheet and not (sheet in sheet_names(real_path) and sheet in sheet_names(generated_path)):
        sheet = None
    rh, rrows = read_table(real_path, sheet=sheet, sku=sku, header=header, delimiter=delimiter)
    gh, grows = read_table(generated_path, sheet=sheet, sku=sku, header=header, delimiter=delimiter)
    real_headerless = False
    if sku and sku not in rh and sku in gh and len(rh) == len(gh):
        # Real file has no header row (e.g. UAMS Workday) but we write one:
        # report the header difference, then align rows by position.
        real_headerless = True
        rrows = read_table(real_path, header=gh, delimiter=delimiter)[1]
    res = {"real_cols": len(rh), "gen_cols": len(gh), "real_rows": len(rrows), "gen_rows": len(grows),
           "header_diffs": [], "framing": None, "sku_missing": False, "only_real": [], "only_gen": [],
           "common": 0, "value_diffs": {}, "type_diffs": {}, "type_diff_skus": {},
           "real_sheets": sheet_names(real_path), "gen_sheets": sheet_names(generated_path),
           "real_headerless": real_headerless}
    if real_headerless:
        res["header_diffs"].append((1, "<no header row>", gh[0] if gh else "<none>"))
        rh = gh
    elif rh != gh:
        for i in range(max(len(rh), len(gh))):
            r = rh[i] if i < len(rh) else "<none>"
            g = gh[i] if i < len(gh) else "<none>"
            if r != g:
                res["header_diffs"].append((i + 1, r, g))

    rf, gf = byte_framing(real_path), byte_framing(generated_path)
    if rf is not None and gf is not None and rf != gf:
        res["framing"] = {"real": rf, "gen": gf}

    sku = sku or rh[0]
    if sku not in rh or sku not in gh:
        res["sku_missing"] = True
        res["ok"] = False
        return res
    ri, gi = rh.index(sku), gh.index(sku)
    rmap = {r[ri][0]: r for r in rrows}
    gmap = {r[gi][0]: r for r in grows}
    res["only_real"] = sorted(set(rmap) - set(gmap))
    res["only_gen"] = sorted(set(gmap) - set(rmap))

    common = set(rmap) & set(gmap)
    res["common"] = len(common)
    # Compare by position (headers can repeat, e.g. two PRICE_AMOUNT columns);
    # when header lists differ, fall back to matching each real column by name.
    if rh == gh:
        pairs = [(i, i, f"{h} [col {i+1}]" if rh.count(h) > 1 else h) for i, h in enumerate(rh)]
    else:
        pairs = [(i, gh.index(h), h) for i, h in enumerate(rh) if h in gh]
    val_diffs, type_diffs = defaultdict(list), defaultdict(lambda: defaultdict(list))
    for k in sorted(common):
        r, g = rmap[k], gmap[k]
        for ri_, gi_, label in pairs:
            rv, rt = r[ri_] if ri_ < len(r) else ("", "blank")
            gv, gt = g[gi_] if gi_ < len(g) else ("", "blank")
            if rv != gv:
                val_diffs[label].append((k, rv, gv))
            if rt != gt:
                type_diffs[label][(rt, gt)].append(k)
    order = [label for _, _, label in pairs]
    res["value_diffs"] = {h: val_diffs[h] for h in order if val_diffs[h]}
    res["type_diffs"] = {h: {t: len(v) for t, v in type_diffs[h].items()} for h in order if type_diffs[h]}
    res["type_diff_skus"] = {h: dict(type_diffs[h]) for h in order if type_diffs[h]}
    res["ok"] = not (res["header_diffs"] or res["framing"] or res["only_real"] or res["only_gen"]
                     or res["value_diffs"] or res["type_diffs"])
    return res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("real")
    ap.add_argument("generated")
    ap.add_argument("--sku", help="SKU header (default: first column)")
    ap.add_argument("--max-diffs", type=int, default=5)
    a = ap.parse_args()

    res = compare(a.real, a.generated, a.sku)
    print(f"real: {res['real_cols']} cols, {res['real_rows']} rows | "
          f"generated: {res['gen_cols']} cols, {res['gen_rows']} rows")
    if res["header_diffs"]:
        print("HEADER MISMATCH")
        for i, r, g in res["header_diffs"]:
            print(f"  col {i}: real={r!r} generated={g!r}")
    if res["framing"]:
        print(f"FRAMING DIFF\n  real={res['framing']['real']}\n  gen ={res['framing']['gen']}")
    if res["sku_missing"]:
        print("SKU column missing from real or generated; cannot align rows")
        return 1
    if res["only_real"] or res["only_gen"]:
        print(f"ROW POPULATION: {len(res['only_real'])} only in real (e.g. {res['only_real'][:5]}), "
              f"{len(res['only_gen'])} only in generated (e.g. {res['only_gen'][:5]})")
    for h in dict.fromkeys(list(res["value_diffs"]) + list(res["type_diffs"])):
        diffs = res["value_diffs"].get(h, [])
        if diffs:
            print(f"VALUE DIFF {h!r}: {len(diffs)}/{res['common']} rows")
            for k, rv, gv in diffs[: a.max_diffs]:
                print(f"    {k}: real={rv!r}\n    {' ' * len(k)}  gen ={gv!r}")
        if h in res["type_diffs"]:
            print(f"TYPE DIFF {h!r}: " + ", ".join(
                f"real {rt} vs gen {gt} x{n}" for (rt, gt), n in res["type_diffs"][h].items()))
    print("MATCH" if res["ok"] else "DIFFERENCES FOUND")
    return 0 if res["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
